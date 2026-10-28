"""
core/attachments.py - turns any number of attached files into model-ready content.

Supports (dozens of files per message):
  - images (png/jpg/webp/bmp/gif)  -> PIL image, auto-shrunk
  - text / code / data files       -> text
  - PDF, audio, short video        -> inline bytes (Gemini reads them natively)
  - Word / Excel / PowerPoint      -> text extracted with python-docx / openpyxl / python-pptx
Budgets keep the request within model limits; whatever is skipped is reported in the
manifest so the model (and the user) know exactly what was and was not read.
"""

from __future__ import annotations

import logging
import os
from pathlib import Path
from typing import Any, Iterable, List, Tuple

logger = logging.getLogger("SaiA")

MAX_FILES = 100
MAX_TEXT_PER_FILE = 60_000
MAX_TEXT_TOTAL = 400_000
MAX_INLINE_PER_FILE = 12 * 1024 * 1024
MAX_INLINE_TOTAL = 18 * 1024 * 1024
MAX_IMAGE_SIDE = 1600

IMAGE_EXT = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".gif"}
TEXT_EXT = {
    ".txt", ".md", ".csv", ".tsv", ".json", ".xml", ".yaml", ".yml", ".ini", ".cfg", ".toml", ".log",
    ".py", ".js", ".ts", ".jsx", ".tsx", ".html", ".htm", ".css", ".scss", ".sql", ".sh", ".bat",
    ".ps1", ".c", ".h", ".cpp", ".hpp", ".cs", ".java", ".kt", ".go", ".rs", ".php", ".rb", ".swift",
    ".r", ".m", ".tex", ".srt", ".vtt", ".env.example",
}
INLINE_MIME = {
    ".pdf": "application/pdf",
    ".wav": "audio/wav", ".mp3": "audio/mp3", ".m4a": "audio/mp4", ".aac": "audio/aac",
    ".ogg": "audio/ogg", ".flac": "audio/flac", ".opus": "audio/ogg",
    ".mp4": "video/mp4", ".mov": "video/quicktime", ".webm": "video/webm", ".mpeg": "video/mpeg",
}
OFFICE_EXT = {".docx", ".xlsx", ".xlsm", ".pptx"}

# one filter string for the file dialog
DIALOG_PATTERN = " ".join(
    "*" + e for e in sorted(IMAGE_EXT | TEXT_EXT | set(INLINE_MIME) | OFFICE_EXT) if e.count(".") == 1
)


def normalize_paths(file_path: Any) -> List[str]:
    """Accepts None, one path, or an iterable of paths; returns unique existing-looking strings."""
    if not file_path:
        return []
    if isinstance(file_path, (str, Path)):
        items: Iterable = [file_path]
    else:
        items = file_path
    seen, result = set(), []
    for item in items:
        p = str(item)
        if p and p not in seen:
            seen.add(p)
            result.append(p)
    return result


def _read_text(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="replace")


def _docx_text(path: Path) -> str:
    import docx  # python-docx

    document = docx.Document(str(path))
    lines = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(lines)


def _xlsx_text(path: Path) -> str:
    import openpyxl

    wb = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
    out: List[str] = []
    for ws in wb.worksheets[:12]:
        out.append(f"## Sheet: {ws.title}")
        for idx, row in enumerate(ws.iter_rows(values_only=True)):
            if idx >= 400:
                out.append("... (rows truncated)")
                break
            if any(v is not None for v in row):
                out.append(" | ".join("" if v is None else str(v) for v in row))
    wb.close()
    return "\n".join(out)


def _pptx_text(path: Path) -> str:
    from pptx import Presentation

    prs = Presentation(str(path))
    out: List[str] = []
    for n, slide in enumerate(prs.slides, 1):
        out.append(f"## Slide {n}")
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
                out.append(shape.text_frame.text.strip())
    return "\n".join(out)


def _clip(text: str, limit: int) -> Tuple[str, bool]:
    return (text[:limit], True) if len(text) > limit else (text, False)


def build_attachment_contents(file_paths: Any) -> Tuple[List[Any], str]:
    """Returns (contents_to_append, manifest_text)."""
    paths = normalize_paths(file_paths)
    skipped_over_limit = max(0, len(paths) - MAX_FILES)
    paths = paths[:MAX_FILES]

    contents: List[Any] = []
    manifest: List[str] = []
    text_used = inline_used = 0

    for n, raw in enumerate(paths, 1):
        p = Path(raw)
        label = f"[{n}] {p.name}"
        if not p.is_file():
            manifest.append(f"{label} - not found, skipped")
            continue
        ext = p.suffix.lower()
        try:
            size = p.stat().st_size
            if ext in IMAGE_EXT:
                from PIL import Image

                with Image.open(p) as img:
                    img.load()
                    copy = img.convert("RGB") if img.mode not in ("RGB", "RGBA", "L") else img.copy()
                copy.thumbnail((MAX_IMAGE_SIDE, MAX_IMAGE_SIDE))
                contents.append(f"\nATTACHMENT {label} (image):")
                contents.append(copy)
                manifest.append(f"{label} - image read")
            elif ext in INLINE_MIME:
                if size > MAX_INLINE_PER_FILE or inline_used + size > MAX_INLINE_TOTAL:
                    manifest.append(f"{label} - too large for this request ({size // 1024} KB), skipped")
                    continue
                from google.genai import types

                contents.append(f"\nATTACHMENT {label} ({INLINE_MIME[ext]}):")
                contents.append(types.Part.from_bytes(data=p.read_bytes(), mime_type=INLINE_MIME[ext]))
                inline_used += size
                kind = "PDF" if ext == ".pdf" else INLINE_MIME[ext].split("/")[0]
                manifest.append(f"{label} - {kind} read")
            elif ext in TEXT_EXT or ext in OFFICE_EXT:
                if text_used >= MAX_TEXT_TOTAL:
                    manifest.append(f"{label} - text budget reached, skipped")
                    continue
                if ext == ".docx":
                    text = _docx_text(p)
                elif ext in (".xlsx", ".xlsm"):
                    text = _xlsx_text(p)
                elif ext == ".pptx":
                    text = _pptx_text(p)
                else:
                    text = _read_text(p)
                text, cut = _clip(text, min(MAX_TEXT_PER_FILE, MAX_TEXT_TOTAL - text_used))
                text_used += len(text)
                contents.append(f"\nATTACHMENT {label} (text):\n{text}")
                manifest.append(f"{label} - text read" + (" (truncated)" if cut else ""))
            else:
                manifest.append(f"{label} - unsupported type '{ext or 'none'}', path: {p}")
        except Exception as exc:  # one bad file must never break the whole request
            logger.warning("attachment %s failed: %s", p, exc)
            manifest.append(f"{label} - could not be read ({type(exc).__name__})")

    if skipped_over_limit:
        manifest.append(f"... {skipped_over_limit} more file(s) ignored (limit {MAX_FILES})")
    text = ""
    if manifest:
        text = "\n\nATTACHMENTS ({} file(s)):\n".format(len(paths)) + "\n".join(manifest) + "\n"
    return contents, text
